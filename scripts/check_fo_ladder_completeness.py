# -*- coding: utf-8 -*-
"""核对副驾驶技术等级阶梯完整性（过程脚本，保留）。

晋升降本：当前最高档为 C 时，777 上应有 FOA1→FOA2→FOB→FOC 首次生效节点。
同机型划转：不强制完整晋升链，另册统计；本脚本仍列出缺失供对照。
"""
from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
import openpyxl.utils

ROOT = Path(__file__).resolve().parents[1]
STRENGTH = Path(r"C:\Users\Administrator\Desktop\技术等级运行资格\南货航飞行实力表20261023R2.xlsx")
TECH = Path(r"C:\Users\Administrator\Desktop\技术等级运行资格\技术等级导出.xlsx")
OUT_XLSX = ROOT / "副驾驶阶梯完整性核对.xlsx"
OUT_TXT = ROOT / "副驾驶阶梯完整性核对.txt"

LADDER = ["FOA1", "FOA2", "FOB", "FOC", "FOD"]
TECH_TO_CODE = {
    "A1类副驾驶": "FOA1",
    "A2类副驾驶": "FOA2",
    "B类副驾驶": "FOB",
    "C类副驾驶": "FOC",
    "D类副驾驶": "FOD",
}


def parse_d(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()[:10]
    for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def bucket(yuan: str) -> str:
    y = yuan or ""
    if "总队777" in y or "上海" in y or "新疆" in y:
        return "同机型划转"
    return "晋升降本"


def highest_from_techinfo(tech: str) -> str | None:
    t = tech or ""
    for name, code in TECH_TO_CODE.items():
        if name in t:
            return code
    return None


def main():
    wb = openpyxl.load_workbook(STRENGTH, read_only=True, data_only=True)
    rows = list(wb["人员信息"].iter_rows(values_only=True))
    wb.close()
    h = list(rows[0])
    idx = {str(x): i for i, x in enumerate(h) if x}
    people = {}
    for r in rows[1:]:
        emp = r[idx["员工号"]]
        if not emp:
            continue
        emp = str(emp)
        tech = str(r[idx["技术信息"]] or "")
        if "副驾驶" not in tech:
            continue
        people[emp] = {
            "姓名": r[idx["姓名"]],
            "分部": r[idx["分部"]],
            "技术信息": tech,
            "原单位": str(r[idx["原单位"]] or ""),
            "册": bucket(str(r[idx["原单位"]] or "")),
            "最高码": highest_from_techinfo(tech),
        }

    wb = openpyxl.load_workbook(TECH, read_only=True, data_only=True)
    tech_rows = list(wb.active.iter_rows(min_row=2, values_only=True))
    wb.close()

    # 777 首次生效：去完全重复后，每人每码最早生效
    by = defaultdict(list)
    seen = set()
    for r in tech_rows:
        emp, name, code, ac = str(r[0]), r[1], r[10], r[13]
        if ac != "777" or code not in LADDER:
            continue
        sd, ed = parse_d(r[14]), parse_d(r[15])
        key = (emp, code, sd, ed)
        if key in seen:
            continue
        seen.add(key)
        if sd:
            by[(emp, code)].append(sd)

    first_node = {}
    for (emp, code), sds in by.items():
        first_node[(emp, code)] = min(sds)

    lines = []
    w = lines.append
    w("副驾驶阶梯完整性核对")
    w("原则：晋升降本副驾驶，清洗后777上应从FOA1到当前最高档级级有首次生效节点。")
    w("同机型划转不强制完整链，本表仍列出缺失供对照。")
    w("")

    detail = []
    for emp, p in sorted(people.items(), key=lambda x: (x[1]["册"], x[1]["最高码"] or "", x[1]["姓名"] or "")):
        hi = p["最高码"]
        if not hi:
            continue
        need = LADDER[: LADDER.index(hi) + 1]
        present = {c: first_node.get((emp, c)) for c in need}
        missing = [c for c in need if present[c] is None]
        ok = len(missing) == 0
        detail.append(
            {
                **p,
                "员工号": emp,
                "应有阶梯": "→".join(need),
                "缺失码": ",".join(missing) if missing else "",
                "是否完整": "是" if ok else "否",
                **{f"{c}_首次生效": present.get(c) for c in LADDER},
            }
        )

    jin = [d for d in detail if d["册"] == "晋升降本"]
    zhuan = [d for d in detail if d["册"] == "同机型划转"]
    jin_bad = [d for d in jin if d["是否完整"] == "否"]
    zhuan_bad = [d for d in zhuan if d["是否完整"] == "否"]

    w(f"实力表副驾驶（能识别最高档）: {len(detail)}")
    w(f"晋升降本: {len(jin)}，阶梯完整 {len(jin)-len(jin_bad)}，不完整 {len(jin_bad)}")
    w(f"同机型划转: {len(zhuan)}，阶梯完整 {len(zhuan)-len(zhuan_bad)}，不完整 {len(zhuan_bad)}（不强制）")
    w("")
    w("=== 晋升降本 · 不完整（点名）===")
    if not jin_bad:
        w("无")
    for d in jin_bad:
        w(
            f"{d['员工号']} {d['姓名']} {d['技术信息']} 原单位={d['原单位']} "
            f"应有={d['应有阶梯']} 缺失={d['缺失码']}"
        )
        for c in LADDER:
            if d.get(f"{c}_首次生效"):
                w(f"  {c}={d[f'{c}_首次生效']}")

    w("")
    w("=== 晋升降本 · 按最高档完整率 ===")
    for code in LADDER:
        sub = [d for d in jin if d["最高码"] == code]
        if not sub:
            continue
        bad = sum(1 for d in sub if d["是否完整"] == "否")
        w(f"{code}: n={len(sub)} 完整={len(sub)-bad} 不完整={bad}")

    OUT_TXT.write_text("\n".join(lines), encoding="utf-8")

    xwb = Workbook()
    ws = xwb.active
    ws.title = "说明"
    for i, t in enumerate(
        [
            "晋升降本：当前最高档之前（含自身）777上FOA1…应级级有首次生效节点。",
            "例：C类副驾驶应有 FOA1、FOA2、FOB、FOC。",
            "同机型划转：不强制；缺失属预期或另册口径。",
            f"晋升降本不完整 {len(jin_bad)} / {len(jin)}；同机型划转不完整 {len(zhuan_bad)} / {len(zhuan)}。",
            f"脚本: {Path(__file__).name}",
        ],
        1,
    ):
        ws[f"A{i}"] = t
    ws.column_dimensions["A"].width = 100

    headers = [
        "员工号",
        "姓名",
        "分部",
        "册",
        "原单位",
        "技术信息",
        "最高码",
        "应有阶梯",
        "是否完整",
        "缺失码",
        "FOA1_首次生效",
        "FOA2_首次生效",
        "FOB_首次生效",
        "FOC_首次生效",
        "FOD_首次生效",
    ]

    def write_sheet(title, data, paint_bad=True):
        ws = xwb.create_sheet(title)
        for c, h in enumerate(headers, 1):
            cell = ws.cell(1, c, h)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9E2F3")
        for i, r in enumerate(data, 2):
            for c, h in enumerate(headers, 1):
                v = r.get(h)
                cell = ws.cell(i, c, v)
                if isinstance(v, date):
                    cell.number_format = "yyyy-mm-dd"
                if paint_bad and h == "是否完整" and v == "否":
                    cell.fill = PatternFill("solid", fgColor="FFC7CE")
        ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(headers))}{max(1, len(data)+1)}"
        ws.freeze_panes = "A2"

    write_sheet("晋升降本", jin)
    write_sheet("晋升降本_不完整", jin_bad)
    write_sheet("同机型划转_对照", zhuan, paint_bad=False)
    xwb.save(OUT_XLSX)
    print("\n".join(lines))
    print("saved", OUT_XLSX)
    print("saved", OUT_TXT)


if __name__ == "__main__":
    main()
