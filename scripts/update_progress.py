"""一二检与副驾驶进度：生成 CSN2API 查询模板 / 离线回填。

本文件仅为辅助脚本，不是业务规则或表格结构的真源。
每次使用前必须先只读调研主人本轮真实原表、查询模板与导出结果，
核对工作表、表头、列位置、分块边界、日期格式和目标单元格。
一切以本轮真实表格及主人已确认的业务口径为准；脚本中的映射仅适用于
已核实的结构。结构变化时先调整映射并核对副本，不能直接套用旧列号或样例。

工作文件和结果默认由调用者指定在仓库根目录，不调用网络接口。
依赖：openpyxl。运行 python scripts/update_progress.py --help。
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import posixpath
import re
import shutil
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl
from openpyxl.utils.datetime import from_excel

ROOT = Path(__file__).resolve().parent.parent
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
SPECS = [("副驾驶进度", "进入时间", "飞行经历"),
         ("一二检进度", "左座进入时间", "左座经历")]


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def save_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def employee_id(value):
    if isinstance(value, bool) or value is None:
        raise ValueError(f"无效员工号：{value!r}")
    if isinstance(value, (int, float)) and int(value) == value:
        value = str(int(value))
    value = str(value).strip()
    if not re.fullmatch(r"\d+", value):
        raise ValueError(f"无效员工号：{value!r}")
    return value


def as_date(value, epoch):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        value = from_excel(value, epoch)
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y%m%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f"无法识别日期：{value!r}")


def integer(value):
    """样表采用去分钟的整数小时；不擅自决定 HH:MM 或小数的取整口径。"""
    if isinstance(value, bool) or value is None or str(value).strip() == "":
        raise ValueError(f"缺失数值：{value!r}")
    try:
        number = Decimal(str(value).strip())
    except InvalidOperation:
        raise ValueError(f"需去分钟版整数，不能自动取整：{value!r}") from None
    if not number.is_finite() or number < 0 or number != number.to_integral_value():
        raise ValueError(f"需非负整数：{value!r}")
    return int(number)


def header_map(ws):
    result = {}
    for cell in ws[1]:
        name = str(cell.value or "").strip()
        if name:
            if name in result:
                raise ValueError(f"{ws.title} 重复列名：{name}")
            result[name] = cell.column
    return result


def roster(source, end, captain_scope="first-block"):
    wb = openpyxl.load_workbook(source, data_only=False)
    rows = []
    try:
        for sheet, date_header, measure in SPECS:
            ws = wb[sheet]
            headers = header_map(ws)
            if ws.cell(1, 1).value != "员工号" or ws.cell(1, 2).value != "姓名":
                raise ValueError(f"{sheet} 员工号/姓名列结构发生变化")
            if date_header not in headers:
                raise ValueError(f"{sheet} 缺少日期列：{date_header}")
            if sheet == "副驾驶进度":
                expected = ("阶段经历时间", "阶段起落数")
            else:
                expected = ("左座经历", "起落数")
            for col, token in zip((6, 7), expected):
                if token not in str(ws.cell(1, col).value):
                    raise ValueError(f"{sheet} F/G 列结构发生变化")
            seen = set()
            for row in range(2, ws.max_row + 1):
                raw_id, name = ws.cell(row, 1).value, ws.cell(row, 2).value
                if raw_id == "员工号":
                    if sheet == "一二检进度" and captain_scope == "first-block":
                        break
                    # 上下块列结构相同，日期仍按同名列；不复制表头。
                    if ws.cell(row, headers[date_header]).value != date_header:
                        raise ValueError(f"{sheet} 第{row}行第二块日期列结构变化")
                    continue
                if raw_id is None and name is None:
                    continue
                rid = employee_id(raw_id)
                name = str(name or "").strip()
                if not name or name.startswith("="):
                    raise ValueError(f"{sheet} 第{row}行姓名缺失或为公式")
                start = as_date(ws.cell(row, headers[date_header]).value, wb.epoch)
                if start > end:
                    raise ValueError(f"{sheet} 第{row}行开始日期晚于结束日期")
                key = (rid, start, end)
                if key in seen:
                    raise ValueError(f"{sheet} 重复人员及日期窗口：{key}")
                seen.add(key)
                rows.append({"sheet":sheet, "row":row, "id":rid, "name":name,
                             "start":start, "end":end, "measure":measure})
    finally:
        wb.close()
    return rows


def preserve_existing(path):
    """覆盖既有产物前保留唯一备份，绝不静默覆盖既有备份。"""
    path = Path(path)
    if path.exists():
        backup = path.with_name(path.stem + "_覆盖前备份_" + datetime.now().strftime("%Y%m%d_%H%M%S_%f") + path.suffix)
        shutil.copy2(path, backup)


def prepare(args):
    end = args.end or datetime.now(timezone(timedelta(hours=8))).date().isoformat()
    end = date.fromisoformat(end).isoformat()
    source, template = Path(args.source).resolve(), Path(args.template).resolve()
    output, session_path = Path(args.output).resolve(), Path(args.session).resolve()
    if output in (source, template) or session_path in (source, template, output):
        raise ValueError("查询产物不能覆盖输入文件")
    source_hash = sha256(source)
    rows = roster(source, end, args.captain_scope)
    wb = openpyxl.load_workbook(template)
    if len(wb.worksheets) != 1 or [wb.active.cell(1, c).value for c in range(1, 5)] != ["员工号", "姓名", "开始日期", "结束日期"]:
        raise ValueError("模板必须为员工号、姓名、开始日期、结束日期四列的单工作表")
    ws = wb.active
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    queries = {}
    for item in rows:
        key = (item["id"], item["start"], item["end"])
        if key in queries and queries[key] != item["name"]:
            raise ValueError(f"员工号与姓名冲突：{key}")
        queries[key] = item["name"]
    for (rid, start, stop), name in queries.items():
        ws.append([rid, name, datetime.fromisoformat(start), datetime.fromisoformat(stop)])
        for col in (3, 4):
            ws.cell(ws.max_row, col).number_format = "yyyy-mm-dd"
    if sha256(source) != source_hash:
        raise ValueError("生成期间原表发生变化，请重新生成")
    preserve_existing(output)
    preserve_existing(session_path)
    wb.save(output)
    wb.close()
    session = {"version":1, "source":str(source), "source_sha256":source_hash,
               "query":str(output), "query_sha256":sha256(output), "end":end,
               "captain_scope":args.captain_scope, "rows":rows,
               "counts":dict(Counter(x["sheet"] for x in rows)), "query_count":len(queries)}
    save_json(session_path, session)
    print(json.dumps({"query":str(output), "session":str(session_path),
                      "end":end, "counts":session["counts"], "queries":len(queries)}, ensure_ascii=False))


def read_export(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    found = []
    required = {"输入员工号", "输入姓名", "输入开始日期", "输入结束日期", "查询状态",
                "员工号", "姓名", "开始日期", "结束日期", "飞行经历", "左座经历", "起落总数"}
    for ws in wb:
        headers = header_map(ws)
        if required <= headers.keys():
            found.append((ws, headers))
    if len(found) != 1:
        raise ValueError("导出中必须有唯一包含完整查询字段的结果表")
    ws, headers = found[0]
    records = defaultdict(list)
    try:
        for line in ws.iter_rows(min_row=2, values_only=True):
            if not any(x is not None for x in line):
                continue
            item = {name:line[col-1] for name,col in headers.items()}
            key = (employee_id(item["输入员工号"]),
                   as_date(item["输入开始日期"], wb.epoch),
                   as_date(item["输入结束日期"], wb.epoch))
            records[key].append(item)
        return records, wb.epoch
    finally:
        wb.close()


def worksheet_parts(zf):
    rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    targets = {r.attrib["Id"]:posixpath.normpath("xl/" + r.attrib["Target"]) if not r.attrib["Target"].startswith("/") else r.attrib["Target"].lstrip("/") for r in rels}
    book = ET.fromstring(zf.read("xl/workbook.xml"))
    return {s.attrib["name"]:targets[s.attrib[f"{{{REL}}}id"]] for s in book.find("s:sheets", NS)}


def patched_sheet(data, changes):
    """仅替换既有数值单元格；保留工作簿其他 XML、缓存、样式等内容。"""
    text = data.decode("utf-8")
    remaining = dict(changes)
    pattern = re.compile(r'<c\b(?P<attrs>[^>]*\br="(?P<ref>[A-Z]+\d+)"[^>]*)(?:>(?P<body>.*?)</c>|\s*/>)', re.S)

    def replace(match):
        ref = match["ref"]
        if ref not in remaining:
            return match[0]
        if re.search(r"<f(?:\s|>)", match["body"] or ""):
            raise ValueError(f"不覆盖目标公式：{ref}")
        value = remaining.pop(ref)
        attrs = re.sub(r'\s+t="[^"]*"', "", match["attrs"]).rstrip(" /\t\r\n")
        return f'<c{attrs}><v>{value}</v></c>'

    text = pattern.sub(replace, text)
    if remaining:
        raise ValueError(f"目标数值单元格不存在：{sorted(remaining)}")
    ET.fromstring(text)
    return text.encode("utf-8")


def report_book(path, summary, details, errors):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "摘要"
    ws.append(["项目", "结果"])
    for k,v in summary.items():
        ws.append([k, str(v)])
    ws = wb.create_sheet("逐项回填")
    ws.append(["工作表", "原行号", "员工号", "姓名", "开始日期", "结束日期", "导出字段", "目标单元格", "原值", "拟写值", "是否变化"])
    for row in details:
        ws.append(row)
    ws = wb.create_sheet("异常")
    ws.append(["说明"])
    for err in errors:
        ws.append([err])
    for ws in wb:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 22
    preserve_existing(path)
    wb.save(path)
    wb.close()


def apply_export(args):
    session = json.loads(Path(args.session).read_text(encoding="utf-8"))
    source = Path(args.source or session["source"]).resolve()
    export = Path(args.export).resolve()
    output = Path(args.output).resolve() if args.output else ROOT / f"{source.stem}（由皮燕子更新副驾驶和一二检数据截止{session['end']}）{source.suffix}"
    report = Path(args.report).resolve() if args.report else None
    protected = {source, export, Path(args.session).resolve(), Path(session["query"]).resolve()}
    if report in protected or report == output or output in protected - {source}:
        raise ValueError("产物路径与输入文件冲突")
    if sha256(source) != session["source_sha256"]:
        raise ValueError("实力表自生成查询表后已变动；请重新核对并生成本轮查询，避免用旧窗口回填")
    rows = roster(source, session["end"], session["captain_scope"])
    if rows != session["rows"]:
        raise ValueError("本轮名单或日期窗口已变化")
    records, epoch = read_export(export)
    requested = {(r["id"], r["start"], r["end"]) for r in rows}
    errors, details, changes = [], [], defaultdict(dict)
    wb = openpyxl.load_workbook(source, data_only=False)
    for item in rows:
        key = (item["id"], item["start"], item["end"])
        matches = records.get(key, [])
        try:
            if len(matches) != 1:
                raise ValueError(f"导出匹配数为{len(matches)}，要求唯一")
            record = matches[0]
            if str(record["查询状态"]).strip() != "成功":
                raise ValueError(f"查询未成功：{record['查询状态']}")
            if employee_id(record["员工号"]) != item["id"]:
                raise ValueError("返回员工号不符")
            for col in ("输入姓名", "姓名"):
                if str(record[col] or "").strip() != item["name"]:
                    raise ValueError(f"{col}不符：{record[col]}")
            for col, expected in (("开始日期", item["start"]), ("结束日期", item["end"])):
                if as_date(record[col], epoch) != expected:
                    raise ValueError(f"返回{col}不符")
            values = [(6, item["measure"], integer(record[item["measure"]])),
                      (7, "起落总数", integer(record["起落总数"]))]
            for col, field, value in values:
                cell = wb[item["sheet"]].cell(item["row"], col)
                if cell.data_type == "f":
                    raise ValueError(f"目标格{cell.coordinate}为公式")
                changed = cell.value != value or cell.data_type != "n"
                details.append([item["sheet"], item["row"], item["id"], item["name"], item["start"], item["end"], field, cell.coordinate, cell.value, value, "是" if changed else "否"])
                if changed:
                    changes[item["sheet"]][cell.coordinate] = value
        except (ValueError, TypeError) as exc:
            errors.append(f"{item['sheet']} 第{item['row']}行 {item['id']} {item['name']}：{exc}")
    wb.close()
    summary = {"模式":"预检" if args.preview else "回填副本" if source != output else "原地回填",
               "原表":str(source), "导出":str(export), "截止日期":session["end"],
               "人员数":len(rows), "核对指标格数":len(details), "拟变化格数":sum(map(len, changes.values())),
               "额外导出窗口数":len(set(records)-requested), "异常数":len(errors)}
    if report:
        report_book(report, summary, details, errors)
    if errors:
        raise ValueError(f"发现{len(errors)}个异常，未生成回填文件：" + "；".join(errors))
    if args.preview:
        print(json.dumps(summary, ensure_ascii=False))
        return
    source_hash = sha256(source)
    # 默认另存最终副本，原表完整保留；原地覆盖才额外创建备份。
    backup = None
    if source == output or args.backup:
        backup = Path(args.backup).resolve() if args.backup else ROOT / (source.stem + "_回填前备份_" + source_hash[:8] + source.suffix)
        if backup in protected or backup in (output, report):
            raise ValueError("备份路径冲突")
        if backup.exists():
            if sha256(backup) != source_hash:
                raise ValueError("既有备份与当前原表不同，请另指定备份路径")
        else:
            shutil.copy2(source, backup)
        if sha256(backup) != source_hash:
            raise ValueError("备份校验失败")
    if output != source and output.exists():
        raise ValueError("最终文件已存在，请核实后替换，不自动生成多份最终版")
    # 先在内存完成全部 XML 变更，异常不落半份输出。
    with ZipFile(source) as zf:
        parts = worksheet_parts(zf)
        patched = {parts[name]:patched_sheet(zf.read(parts[name]), cells) for name,cells in changes.items()}
        entries = [(info, patched.get(info.filename, zf.read(info.filename))) for info in zf.infolist()]
        archive_comment = zf.comment
    if sha256(source) != source_hash:
        raise ValueError("回填期间原表发生变化，停止写入")
    staged = output.with_name(output.stem + "_写入中" + output.suffix)
    if staged in protected or staged in (backup, report):
        raise ValueError("中间文件路径冲突")
    preserve_existing(staged)
    with ZipFile(staged, "w") as zf:
        zf.comment = archive_comment
        for info, data in entries:
            zf.writestr(info, data)
    result = openpyxl.load_workbook(staged, data_only=False)
    try:
        for row in details:
            if result[row[0]][row[7]].value != row[9]:
                raise ValueError(f"回填校验失败：{row[0]} {row[7]}")
    finally:
        result.close()
    if sha256(source) != source_hash:
        raise ValueError("写入前原表已变动；中间副本保留，停止替换")
    os.replace(staged, output)
    print(json.dumps({**summary, "输出":str(output), "原表保留":str(source), "备份":str(backup) if backup else None, "报告":str(report) if report else None}, ensure_ascii=False))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    subs = parser.add_subparsers(dest="command", required=True)
    p = subs.add_parser("prepare", help="生成查询模板和本轮记录")
    p.add_argument("--source", required=True)
    p.add_argument("--template", required=True)
    p.add_argument("--end", help="上海日期 YYYY-MM-DD；Agent 调用应明确传入")
    p.add_argument("--captain-scope", choices=("first-block", "all"), default="first-block")
    p.add_argument("--output", required=True)
    p.add_argument("--session", required=True)
    p.set_defaults(run=prepare)
    p = subs.add_parser("apply", help="离线核对导出并回填；--preview 仅预检")
    p.add_argument("--source", help="默认使用本轮记录中的原表")
    p.add_argument("--session", required=True)
    p.add_argument("--export", required=True)
    p.add_argument("--output", help="默认按原表名和本轮截止日期生成唯一最终版文件名")
    p.add_argument("--report", help="可选：仅在需要过程报告时生成")
    p.add_argument("--backup")
    p.add_argument("--preview", action="store_true")
    p.set_defaults(run=apply_export)
    args = parser.parse_args()
    try:
        args.run(args)
    except (ValueError, KeyError, OSError) as exc:
        parser.exit(1, f"停止：{exc}\n")


if __name__ == "__main__":
    main()
