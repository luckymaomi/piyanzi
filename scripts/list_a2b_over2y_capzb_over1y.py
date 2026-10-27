# -*- coding: utf-8 -*-
"""货航：副驾驶 A2→B 超 2 年名单；机长 CAPZ→CAPB 超 1 年名单。Excel×2 + Markdown 说明。"""
from __future__ import annotations

import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(r"C:\Users\Administrator\Desktop\技术等级运行资格")
SP = BASE / "南货航飞行实力表20261023R2.xlsx"
TG = BASE / "技术等级导出.xlsx"
ROOT = Path(__file__).resolve().parents[1]
CUTOFF = date(2026, 10, 23)
FO_DAYS = 365 * 2  # 2 年
CAP_DAYS = 365  # 1 年
OUT_FO = ROOT / "副驾驶_A2转B超过2年名单_20261023.xlsx"
OUT_CAP = ROOT / "机长_CAPZ到CAPB超过1年名单_20261023.xlsx"
OUT_MD = ROOT / "超期名单_A2转B与CAPZ到CAPB_说明_20261023.md"


def to_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()[:10]
    for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def load(path, sheet=None):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    h = [str(x).strip() if x is not None else "" for x in rows[0]]
    return {x: i for i, x in enumerate(h) if x}, rows[1:]


def yuan_bucket(yuan: str) -> str:
    y = yuan or ""
    if "返聘" in y:
        return "返聘"
    if any(k in y for k in ("总队777", "上海", "新疆")):
        return "同机型划转"
    return "晋升降本"


def role_of(tech: str) -> str:
    t = tech or ""
    if "副驾驶" in t:
        return "副驾驶"
    if "飞行教员" in t or "教员" in t:
        return "教员"
    if "机长" in t:
        return "机长"
    return "其他"


def style_header(ws, n):
    fill = PatternFill("solid", fgColor="D9E2F3")
    for c in range(1, n + 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(n)}{ws.max_row}"


def main():
    import openpyxl.utils

    pi, prow = load(SP, "人员信息")
    ti, trow = load(TG)

    people = {}
    for r in prow:
        emp = str(r[pi["员工号"]] or "").strip()
        if not emp:
            continue
        tech = str(r[pi["技术信息"]] or "")
        people[emp] = dict(
            员工号=emp,
            姓名=r[pi["姓名"]],
            分部=r[pi["分部"]],
            技术信息=tech,
            原单位=r[pi["原单位"]],
            角色=role_of(tech),
            册=yuan_bucket(str(r[pi["原单位"]] or "")),
        )

    need_codes = {"FOA2", "FOB", "CAPZ", "CAPB"}
    first = defaultdict(dict)  # emp -> code -> date
    name_from_tg = {}
    for r in trow:
        emp = str(r[ti["员工号"]] or "").strip()
        if not emp:
            continue
        code = str(r[ti["技术等级代码"]] or "").strip().upper()
        if code not in need_codes:
            continue
        if str(r[ti["机型"]] or "").strip() != "777":
            continue
        d = to_date(r[ti["生效日期"]])
        if not d:
            continue
        name_from_tg[emp] = r[ti["姓名"]]
        prev = first[emp].get(code)
        if prev is None or d < prev:
            first[emp][code] = d

    def person(emp):
        if emp in people:
            return people[emp]
        return dict(
            员工号=emp,
            姓名=name_from_tg.get(emp, ""),
            分部="",
            技术信息="",
            原单位="",
            角色="",
            册="",
        )

    # —— 副驾驶 A2→B ——
    fo_done = []
    fo_stuck = []
    fo_pool = [e for e, p in people.items() if p["角色"] == "副驾驶"]
    for emp in fo_pool:
        a2 = first[emp].get("FOA2")
        b = first[emp].get("FOB")
        p = person(emp)
        if a2 and b and b >= a2:
            days = (b - a2).days
            if days > FO_DAYS:
                fo_done.append({**p, "FOA2首次生效": a2, "FOB首次生效": b, "间隔天": days, "类型": "已完成且间隔>2年"})
        elif a2 and not b:
            days = (CUTOFF - a2).days
            if days > FO_DAYS:
                fo_stuck.append(
                    {**p, "FOA2首次生效": a2, "FOB首次生效": None, "已耗天": days, "类型": "仍无B且自A2已超2年"}
                )

    # —— 机长 CAPZ→CAPB（现势机长/教员；节点在 777）——
    cap_done = []
    cap_stuck = []
    cap_pool = [e for e, p in people.items() if p["角色"] in ("机长", "教员")]
    # 也纳入有 CAPZ 节点但不在人员信息机长侧的？按实力表现势机长/教员为主
    for emp in cap_pool:
        z = first[emp].get("CAPZ")
        b = first[emp].get("CAPB")
        p = person(emp)
        if z and b and b >= z:
            days = (b - z).days
            if days > CAP_DAYS:
                cap_done.append({**p, "CAPZ首次生效": z, "CAPB首次生效": b, "间隔天": days, "类型": "已完成且间隔>1年"})
        elif z and not b:
            days = (CUTOFF - z).days
            if days > CAP_DAYS:
                cap_stuck.append(
                    {**p, "CAPZ首次生效": z, "CAPB首次生效": None, "已耗天": days, "类型": "仍无CAPB且自CAPZ已超1年"}
                )

    # Excel FO
    wb = Workbook()
    ws = wb.active
    ws.title = "00_口径"
    ws.append(["项", "内容"])
    for row in [
        ("任务", "副驾驶 A2转B 超过 2 年名单"),
        ("实力表", str(SP)),
        ("技术等级", str(TG)),
        ("截止日", str(CUTOFF)),
        ("机型", "777"),
        ("节点", "FOA2 / FOB 同码取首次生效日"),
        ("名单范围", "实力表人员信息·技术信息含副驾驶"),
        ("阈值1", "已完成：FOB生效−FOA2生效 > 730 天"),
        ("阈值2", "在途：有FOA2无FOB，且截止日−FOA2 > 730 天"),
        ("分册", "册=原单位晋升降本/同机型划转/返聘，仅标注不剔除"),
    ]:
        ws.append(list(row))
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 90

    def write_sheet(wb, title, rows, headers):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for r in sorted(rows, key=lambda x: (-(x.get("间隔天") or x.get("已耗天") or 0), x["员工号"])):
            ws.append([r.get(h) for h in headers])
        if rows:
            style_header(ws, len(headers))
        return ws

    h_done = ["员工号", "姓名", "分部", "册", "技术信息", "原单位", "FOA2首次生效", "FOB首次生效", "间隔天", "类型"]
    h_stuck = ["员工号", "姓名", "分部", "册", "技术信息", "原单位", "FOA2首次生效", "已耗天", "类型"]
    write_sheet(wb, "01_已完成间隔超2年", fo_done, h_done)
    write_sheet(wb, "02_在途超2年未到B", fo_stuck, h_stuck)
    ws = wb.create_sheet("03_汇总")
    ws.append(["类别", "人数"])
    ws.append(["已完成且间隔>2年", len(fo_done)])
    ws.append(["在途超2年未到B", len(fo_stuck)])
    ws.append(["副驾驶池人数", len(fo_pool)])
    wb.save(OUT_FO)

    # Excel CAP
    wb = Workbook()
    ws = wb.active
    ws.title = "00_口径"
    ws.append(["项", "内容"])
    for row in [
        ("任务", "机长 CAPZ 到 CAPB 超过 1 年名单"),
        ("实力表", str(SP)),
        ("技术等级", str(TG)),
        ("截止日", str(CUTOFF)),
        ("机型", "777"),
        ("节点", "CAPZ / CAPB 同码取首次生效日"),
        ("名单范围", "实力表人员信息·技术信息为机长或教员"),
        ("阈值1", "已完成：CAPB生效−CAPZ生效 > 365 天"),
        ("规则2", "在途：有CAPZ无CAPB，且截止日−CAPZ > 365 天"),
    ]:
        ws.append(list(row))
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 90
    h_done = ["员工号", "姓名", "分部", "册", "技术信息", "原单位", "CAPZ首次生效", "CAPB首次生效", "间隔天", "类型"]
    h_stuck = ["员工号", "姓名", "分部", "册", "技术信息", "原单位", "CAPZ首次生效", "已耗天", "类型"]
    write_sheet(wb, "01_已完成间隔超1年", cap_done, h_done)
    write_sheet(wb, "02_在途超1年未到CAPB", cap_stuck, h_stuck)
    ws = wb.create_sheet("03_汇总")
    ws.append(["类别", "人数"])
    ws.append(["已完成且间隔>1年", len(cap_done)])
    ws.append(["在途超1年未到CAPB", len(cap_stuck)])
    ws.append(["机长+教员池人数", len(cap_pool)])
    wb.save(OUT_CAP)

    md = f"""# 超期名单说明：A2→B（>2年）与 CAPZ→CAPB（>1年）

## 0. 本轮做了什么

按南货航飞行部潜规则：齐表 → 777 同码首次生效清洗 → 出超期名单。  
交付：两个 Excel + 本说明。

| 项 | 值 |
| --- | --- |
| 截止日 | {CUTOFF} |
| 实力表 | `{SP}` |
| 技术等级 | `{TG}` |

---

## 1. 清洗后是什么样

1. 机型只留 **777**；资格码取**首次生效日**。  
2. 副驾驶池：人员信息技术信息含「副驾驶」→ **{len(fo_pool)}** 人。  
3. 机长侧池：技术信息为机长或教员 → **{len(cap_pool)}** 人。  
4. 原单位册（晋升降本 / 同机型划转 / 返聘）仅标注，本轮**不按册剔除**（主人未要求只做晋升降本）。

---

## 2. 按什么口径统计了什么

### 2.1 副驾驶 A2 转 B 超过 2 年

1. 起点：FOA2 首次生效；终点：FOB 首次生效。  
2. **已完成超期**：两端都有，且间隔 **> 730 天** → **{len(fo_done)}** 人。  
3. **在途超期**：有 FOA2、尚无 FOB，截止日已耗 **> 730 天** → **{len(fo_stuck)}** 人。  

附件：`{OUT_FO.name}`（sheet 01 / 02）。

### 2.2 机长 CAPZ 到 CAPB 超过 1 年

1. 起点：CAPZ 首次生效；终点：CAPB 首次生效。  
2. **已完成超期**：间隔 **> 365 天** → **{len(cap_done)}** 人。  
3. **在途超期**：有 CAPZ、尚无 CAPB，已耗 **> 365 天** → **{len(cap_stuck)}** 人。  

附件：`{OUT_CAP.name}`（sheet 01 / 02）。

---

## 3. 结论（可回业务）

1. 副驾驶 A2→B：已完成间隔超 2 年 **{len(fo_done)}** 人；仍停在 A2 且已超 2 年 **{len(fo_stuck)}** 人。详见 `{OUT_FO.name}`。  
2. 机长 CAPZ→CAPB：已完成间隔超 1 年 **{len(cap_done)}** 人；仍停在 CAPZ 且已超 1 年 **{len(cap_stuck)}** 人。详见 `{OUT_CAP.name}`。  
3. 若只要「已经转完且耗时长」或只要「还没转到、已经拖超时」，用对应 sheet；两盘都点了名，避免漏「卡在半路」的人。

---

## 4. 附件

1. `{OUT_FO}`  
2. `{OUT_CAP}`  
3. 脚本：`scripts/list_a2b_over2y_capzb_over1y.py`  
"""
    OUT_MD.write_text(md, encoding="utf-8")
    print(md)
    print("FO done", len(fo_done), "stuck", len(fo_stuck))
    print("CAP done", len(cap_done), "stuck", len(cap_stuck))


if __name__ == "__main__":
    main()
