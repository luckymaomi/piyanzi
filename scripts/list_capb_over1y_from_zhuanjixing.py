# -*- coding: utf-8 -*-
"""一二检进度：转机型开始时间 → CAPB，超过1年名单。Excel + Markdown。"""
from __future__ import annotations

import sys
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
import openpyxl.utils

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(r"C:\Users\Administrator\Desktop\技术等级运行资格")
SP = BASE / "南货航飞行实力表20261023R2.xlsx"
TG = BASE / "技术等级导出.xlsx"
ROOT = Path(__file__).resolve().parents[1]
CUTOFF = date(2026, 10, 23)
LIMIT = 365
OUT_X = ROOT / "机长_转机型开始到CAPB超过1年名单_20261023.xlsx"
OUT_MD = ROOT / "机长_转机型开始到CAPB超过1年_说明_20261023.md"


def to_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)):
        try:
            return (datetime(1899, 12, 30) + timedelta(days=float(v))).date()
        except Exception:
            return None
    s = str(v).strip()[:10]
    for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def load(path, sheet=None):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    h = [str(x).strip().replace("\n", "") if x is not None else "" for x in rows[0]]
    return {x: i for i, x in enumerate(h) if x}, rows[1:], h


def style_header(ws, n):
    fill = PatternFill("solid", fgColor="D9E2F3")
    for c in range(1, n + 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    if ws.max_row >= 1:
        ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(n)}{max(ws.max_row, 1)}"


def main():
    yi, yrows, _ = load(SP, "一二检进度")
    ti, trows, _ = load(TG)
    pi, prows, _ = load(SP, "人员信息")

    # 人员信息补充
    info = {}
    for r in prows:
        emp = str(r[pi["员工号"]] or "").strip()
        if not emp:
            continue
        info[emp] = dict(
            技术信息=r[pi["技术信息"]],
            原单位=r[pi["原单位"]],
            分部=r[pi["分部"]],
        )

    # CAPB 首次生效 777
    capb = {}
    for r in trows:
        emp = str(r[ti["员工号"]] or "").strip()
        if not emp:
            continue
        if str(r[ti["技术等级代码"]] or "").strip().upper() != "CAPB":
            continue
        if str(r[ti["机型"]] or "").strip() != "777":
            continue
        d = to_date(r[ti["生效日期"]])
        if not d:
            continue
        if emp not in capb or d < capb[emp]:
            capb[emp] = d

    progress = []
    for r in yrows:
        emp = str(r[yi["员工号"]] or "").strip()
        if not emp:
            continue
        start = to_date(r[yi["转机型开始时间"]])
        progress.append(
            dict(
                员工号=emp,
                姓名=r[yi["姓名"]],
                转机型开始时间=start,
                阶段=r[yi["阶段"]] if "阶段" in yi else None,
                左座进入时间=to_date(r[yi["左座进入时间"]]) if "左座进入时间" in yi else None,
                首次一检=to_date(r[yi["首次一检"]]) if "首次一检" in yi else None,
                首次二检=r[yi["首次二检"]] if "首次二检" in yi else None,
                报务资格=r[yi["报务资格"]] if "报务资格" in yi else None,
            )
        )

    done = []
    stuck = []
    early_capb = []  # CAPB 早于转机型开始
    no_start = []
    ok_within = 0

    for p in progress:
        emp = p["员工号"]
        start = p["转机型开始时间"]
        b = capb.get(emp)
        extra = info.get(emp, {})
        base = {
            **p,
            "技术信息": extra.get("技术信息"),
            "原单位": extra.get("原单位"),
            "分部_人员信息": extra.get("分部"),
            "CAPB首次生效": b,
        }
        if not start:
            no_start.append(base)
            continue
        if b and b < start:
            early_capb.append({**base, "间隔天": (b - start).days, "类型": "CAPB早于转机型开始"})
            continue
        if b and b >= start:
            days = (b - start).days
            if days > LIMIT:
                done.append({**base, "间隔天": days, "类型": "已获CAPB且间隔>1年"})
            else:
                ok_within += 1
        else:
            days = (CUTOFF - start).days
            if days > LIMIT:
                stuck.append({**base, "已耗天": days, "类型": "尚无CAPB且自转机型开始已超1年"})

    wb = Workbook()
    ws = wb.active
    ws.title = "00_口径"
    for a, b in [
        ("任务", "一二检进度：转机型开始时间 → CAPB，超过1年"),
        ("实力表", str(SP)),
        ("一二检进度人数", len(progress)),
        ("技术等级", str(TG)),
        ("CAPB", "777 首次生效日"),
        ("截止日", str(CUTOFF)),
        ("超期阈值", f"> {LIMIT} 天"),
        ("已完成超期", len(done)),
        ("在途超期", len(stuck)),
        ("CAPB早于转机型开始", len(early_capb)),
        ("无转机型开始时间", len(no_start)),
        ("已获CAPB且间隔≤1年", ok_within),
    ]:
        ws.append([a, b])
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 90

    def sheet(name, rows, headers):
        w = wb.create_sheet(name)
        w.append(headers)
        for r in sorted(rows, key=lambda x: (-(x.get("间隔天") or x.get("已耗天") or 0), x["员工号"])):
            w.append([r.get(h) for h in headers])
        if rows:
            style_header(w, len(headers))

    h_done = [
        "员工号",
        "姓名",
        "转机型开始时间",
        "CAPB首次生效",
        "间隔天",
        "阶段",
        "技术信息",
        "原单位",
        "类型",
    ]
    h_stuck = [
        "员工号",
        "姓名",
        "转机型开始时间",
        "已耗天",
        "阶段",
        "技术信息",
        "原单位",
        "左座进入时间",
        "首次一检",
        "报务资格",
        "类型",
    ]
    sheet("01_已获CAPB间隔超1年", done, h_done)
    sheet("02_尚无CAPB已超1年", stuck, h_stuck)
    sheet(
        "03_CAPB早于转机型开始",
        early_capb,
        ["员工号", "姓名", "转机型开始时间", "CAPB首次生效", "间隔天", "阶段", "技术信息", "原单位", "类型"],
    )
    ws = wb.create_sheet("04_汇总")
    ws.append(["类别", "人数"])
    for k, v in [
        ("一二检进度在册", len(progress)),
        ("已获CAPB且间隔>1年", len(done)),
        ("尚无CAPB且已超1年", len(stuck)),
        ("CAPB早于转机型开始（另列）", len(early_capb)),
        ("无转机型开始时间", len(no_start)),
        ("已获CAPB且间隔≤1年", ok_within),
    ]:
        ws.append([k, v])

    wb.save(OUT_X)

    md = f"""# 转机型开始时间 → CAPB 超过 1 年

## 0. 本轮做了什么

按主人要求重算：起点改用实力表「一二检进度」**转机型开始时间**，终点为技术等级 **CAPB（777）首次生效**；超过 1 年（>365 天）出名单。

| 项 | 值 |
| --- | --- |
| 截止日 | {CUTOFF} |
| 实力表 | `{SP}` · sheet 一二检进度 |
| 技术等级 | `{TG}` |

---

## 1. 清洗后是什么样

| 项 | 数量 |
| --- | ---: |
| 一二检进度在册 | {len(progress)} |
| 其中能对上 CAPB 节点（任意时间） | {sum(1 for p in progress if p['员工号'] in capb)} |
| 无转机型开始时间 | {len(no_start)} |

---

## 2. 按什么口径统计了什么

1. **起点**：一二检进度 · 转机型开始时间。  
2. **终点**：777 CAPB 首次生效日。  
3. **已完成超期**：已有 CAPB，且 CAPB ≥ 转机型开始，间隔 **> 365 天** → **{len(done)}** 人。  
4. **在途超期**：尚无 CAPB（或 CAPB 未取到），截止日 − 转机型开始 **> 365 天** → **{len(stuck)}** 人。  
5. CAPB 早于转机型开始 → 另列 {len(early_capb)} 人，不进上面两盘。

---

## 3. 结论

1. 已获 CAPB 且自转机型开始超过 1 年：**{len(done)}** 人（见 Excel 01）。  
2. 尚无 CAPB 且自转机型开始已超过 1 年：**{len(stuck)}** 人（见 Excel 02）。  
3. 与旧口径「CAPZ→CAPB」不同：本轮**不以 CAPZ 为起点**，以一二检进度转机型开始时间为准。

---

## 4. 附件

1. `{OUT_X}`  
2. 脚本：`scripts/list_capb_over1y_from_zhuanjixing.py`  
"""
    OUT_MD.write_text(md, encoding="utf-8")
    print(md)


if __name__ == "__main__":
    main()
