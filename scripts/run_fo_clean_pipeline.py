# -*- coding: utf-8 -*-
"""按当前口径跑一遍副驾驶技术等级清洗，并做合理性检查（过程脚本，保留）。"""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
import openpyxl.utils

ROOT = Path(__file__).resolve().parents[1]
STRENGTH = Path(r"C:\Users\Administrator\Desktop\技术等级运行资格\南货航飞行实力表20261023R2.xlsx")
TECH = Path(r"C:\Users\Administrator\Desktop\技术等级运行资格\技术等级导出.xlsx")
OUT_XLSX = ROOT / "副驾驶清洗实跑结果.xlsx"
OUT_TXT = ROOT / "副驾驶清洗实跑结果.txt"

LADDER = ["FOA1", "FOA2", "FOB", "FOC", "FOD"]
TECH_TO_CODE = {
    "A1类副驾驶": "FOA1",
    "A2类副驾驶": "FOA2",
    "B类副驾驶": "FOB",
    "C类副驾驶": "FOC",
    "D类副驾驶": "FOD",
}


def parse_d(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()[:10]
    for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def bucket(yuan: str) -> str:
    y = yuan or ""
    if "总队777" in y or "上海" in y or "新疆" in y:
        return "同机型划转"
    return "晋升降本"


def highest_from_techinfo(tech: str):
    for name, code in TECH_TO_CODE.items():
        if name in (tech or ""):
            return code
    return None


def main():
    lines = []
    w = lines.append

    # --- load people ---
    wb = openpyxl.load_workbook(STRENGTH, read_only=True, data_only=True)
    rows = list(wb["人员信息"].iter_rows(values_only=True))
    wb.close()
    h = list(rows[0])
    idx = {str(x): i for i, x in enumerate(h) if x}
    people = {}
    for r in rows[1:]:
        emp = r[idx["员工号"]]
        if not emp:
            continue
        emp = str(emp)
        tech = str(r[idx["技术信息"]] or "")
        if "副驾驶" not in tech:
            continue
        people[emp] = {
            "姓名": r[idx["姓名"]],
            "分部": r[idx["分部"]],
            "技术信息": tech,
            "原单位": str(r[idx["原单位"]] or ""),
            "册": bucket(str(r[idx["原单位"]] or "")),
            "最高码": highest_from_techinfo(tech),
        }

    # --- load tech long ---
    wb = openpyxl.load_workbook(TECH, read_only=True, data_only=True)
    tech_raw = list(wb.active.iter_rows(min_row=2, values_only=True))
    wb.close()

    stats = Counter()
    stats["技术等级原始行"] = len(tech_raw)

    # Step1: 777 + FO ladder only (for FO pipeline)
    step1 = []
    for r in tech_raw:
        code, ac = r[10], r[13]
        if code not in LADDER:
            continue
        if ac != "777":
            stats["洗掉_非777阶梯段"] += 1
            continue
        step1.append(r)
    stats["留_777阶梯段"] = len(step1)

    # Step2: exact dedupe
    seen = set()
    step2 = []
    for r in step1:
        emp, code, sd, ed = str(r[0]), r[10], parse_d(r[14]), parse_d(r[15])
        k = (emp, code, sd, ed)
        if k in seen:
            stats["洗掉_完全重复"] += 1
            continue
        seen.add(k)
        if not sd:
            stats["洗掉_缺生效日"] += 1
            continue
        step2.append({"emp": emp, "name": r[1], "code": code, "sd": sd, "ed": parse_d(r[15])})
    stats["留_去重后阶梯段"] = len(step2)

    # first node per emp+code
    by = defaultdict(list)
    for r in step2:
        by[(r["emp"], r["code"])].append(r)
    nodes = {}
    for key, items in by.items():
        items = sorted(items, key=lambda x: x["sd"])
        nodes[key] = items[0]
        if len(items) > 1:
            stats["洗掉_同码续期非首次段"] += len(items) - 1

    # Step3: split cohort + build person table
    result_rows = []
    for emp, p in people.items():
        hi = p["最高码"]
        if not hi:
            stats["跳过_技术信息无副驾档"] += 1
            continue
        need = LADDER[: LADDER.index(hi) + 1]
        present = {c: nodes.get((emp, c)) for c in need}
        missing = [c for c in need if present[c] is None]
        foa1 = nodes.get((emp, "FOA1"))
        hi_node = nodes.get((emp, hi))
        days = None
        if p["册"] == "晋升降本" and foa1 and hi_node and hi != "FOA1":
            days = (hi_node["sd"] - foa1["sd"]).days
        result_rows.append(
            {
                "员工号": emp,
                "姓名": p["姓名"],
                "分部": p["分部"],
                "册": p["册"],
                "原单位": p["原单位"],
                "技术信息": p["技术信息"],
                "最高码": hi,
                "应有阶梯": "→".join(need),
                "阶梯完整": "是" if not missing else "否",
                "缺失码": ",".join(missing),
                "FOA1首次": foa1["sd"] if foa1 else None,
                "最高档首次": hi_node["sd"] if hi_node else None,
                "运行到最高档天数_仅晋升降本": days,
                **{f"{c}_首次": (present[c]["sd"] if c in present and present[c] else None) for c in LADDER},
            }
        )

    jin = [r for r in result_rows if r["册"] == "晋升降本"]
    zhuan = [r for r in result_rows if r["册"] == "同机型划转"]
    jin_bad = [r for r in jin if r["阶梯完整"] == "否"]
    jin_cyc = [r for r in jin if r["运行到最高档天数_仅晋升降本"] is not None]

    w("========== 清洗规则（逻辑顺序）==========")
    w("1. 机型只留777（非777阶梯段洗掉）")
    w("2. 完全重复去重；无生效日丢掉")
    w("3. 同人同码多段：只留最早生效=首次节点，更晚续期不当首次")
    w("4. 按实力表原单位分册：总队777/上海/新疆=同机型划转；其余=晋升降本")
    w("5. 晋升降本：必须FOA1→当前最高档级级有首次节点；周期=FOA1首次→最高档首次")
    w("6. 同机型划转：不进晋升降本周期均值；不因缺2021后FOA1当异常")
    w("")
    w("========== 本轮实跑量 ==========")
    for k, v in stats.items():
        w(f"{k}: {v}")
    w("")
    w(f"实力表副驾驶入样: {len(result_rows)}")
    w(f"同机型划转: {len(zhuan)}")
    w(f"晋升降本: {len(jin)}")
    w(f"晋升降本阶梯完整: {len(jin)-len(jin_bad)} / {len(jin)}")
    w(f"晋升降本阶梯不完整: {len(jin_bad)}")
    if jin_bad:
        for r in jin_bad:
            w(f"  异常 {r['员工号']} {r['姓名']} 缺 {r['缺失码']} {r['技术信息']} 原单位={r['原单位']}")
    else:
        w("  （无异常点名）")
    w("")
    w("========== 晋升降本 · 运行到最高档（天）合理性 ==========")
    by_hi = defaultdict(list)
    for r in jin_cyc:
        by_hi[r["最高码"]].append(r["运行到最高档天数_仅晋升降本"])
    for code in LADDER[1:]:
        arr = sorted(by_hi.get(code, []))
        if not arr:
            w(f"{code}: 无样本")
            continue
        avg = sum(arr) / len(arr)
        w(f"{code}: n={len(arr)} min={arr[0]} max={arr[-1]} avg={avg:.1f} med={arr[len(arr)//2]}")
        # name extremes
        sub = [r for r in jin_cyc if r["最高码"] == code]
        lo = min(sub, key=lambda x: x["运行到最高档天数_仅晋升降本"])
        hi = max(sub, key=lambda x: x["运行到最高档天数_仅晋升降本"])
        w(f"  最快 {lo['姓名']}({lo['员工号']}) {lo['运行到最高档天数_仅晋升降本']}天 FOA1={lo['FOA1首次']} → {lo['最高档首次']}")
        w(f"  最慢 {hi['姓名']}({hi['员工号']}) {hi['运行到最高档天数_仅晋升降本']}天 FOA1={hi['FOA1首次']} → {hi['最高档首次']}")

    # 彭田
    pt = next((r for r in jin if r["姓名"] == "彭田"), None)
    w("")
    w("========== 点名抽查 ==========")
    if pt:
        w(
            f"彭田({pt['员工号']}) 册={pt['册']} 完整={pt['阶梯完整']} "
            f"FOA1={pt['FOA1首次']} FOD={pt['FOD_首次']} 天={pt['运行到最高档天数_仅晋升降本']}"
        )
    # one 划转 D
    wr = next((r for r in zhuan if r["最高码"] == "FOD" and r["姓名"] == "黄俊豪"), None)
    if wr:
        w(
            f"黄俊豪({wr['员工号']}) 册={wr['册']} 完整={wr['阶梯完整']} "
            f"FOA1={wr['FOA1首次']} FOD={wr['FOD_首次']} （划转册不进晋升均值）"
        )

    w("")
    reasonable = len(jin_bad) == 0 and len(jin) > 0
    w("========== 结论 ==========")
    if reasonable:
        w("合理：晋升降本阶梯全完整；划转已分册；非777/重复/续期非首次已按规则处理。")
    else:
        w("需人工：晋升降本存在阶梯缺口，见上方点名。")

    OUT_TXT.write_text("\n".join(lines), encoding="utf-8")

    xwb = Workbook()
    ws = xwb.active
    ws.title = "清洗说明与结论"
    for i, t in enumerate(lines, 1):
        ws[f"A{i}"] = t
    ws.column_dimensions["A"].width = 110

    headers = list(result_rows[0].keys()) if result_rows else []

    def sheet(title, data):
        ws = xwb.create_sheet(title)
        for c, h in enumerate(headers, 1):
            cell = ws.cell(1, c, h)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9E2F3")
        for i, r in enumerate(data, 2):
            for c, h in enumerate(headers, 1):
                v = r.get(h)
                cell = ws.cell(i, c, v)
                if isinstance(v, date):
                    cell.number_format = "yyyy-mm-dd"
                if h == "阶梯完整" and v == "否":
                    cell.fill = PatternFill("solid", fgColor="FFC7CE")
        ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(headers))}{max(1,len(data)+1)}"
        ws.freeze_panes = "A2"

    sheet("全量结果", result_rows)
    sheet("晋升降本", jin)
    sheet("同机型划转", zhuan)
    xwb.save(OUT_XLSX)
    print("\n".join(lines))
    print("saved", OUT_XLSX)


if __name__ == "__main__":
    main()
