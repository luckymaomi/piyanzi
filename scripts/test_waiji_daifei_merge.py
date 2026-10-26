"""外籍批复核心语义回归；直接运行，无额外测试框架依赖。"""
import sys
sys.dont_write_bytecode = True
import tempfile
import unittest
from contextlib import redirect_stderr
from io import StringIO
from datetime import date, datetime
from pathlib import Path

import openpyxl
import waiji_daifei_merge as app


class ApprovalTests(unittest.TestCase):
    def test_mexico_exact_airports_and_city_expansion(self):
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            paths = [directory / 'open.xlsx', directory / 'temporary.xlsx']
            for path in paths:
                workbook = openpyxl.Workbook()
                for i, station in enumerate(['墨西哥城空运机场', '墨西哥城国际', '墨西哥城'], 1):
                    workbook.active.append([i, 'CSG', '广州', station, 'G', 'B', '否'])
                workbook.save(path)
            output = directory / 'output.xlsx'
            report = app.run_merge(*paths, app.DEFAULT_TEMPLATE, output, date(2026, 10, 25), date(2027, 3, 27))
            self.assertEqual(report['unknown_airports'], [])
            workbook = openpyxl.load_workbook(output)
            ports = dict(workbook[app.SHEET_APT].iter_rows(min_row=2, max_col=2, values_only=True))
            self.assertEqual(ports['墨西哥城空运机场'], 'NLU')
            self.assertEqual(ports['墨西哥城国际'], 'MEX')
            names = ['墨西哥城空运机场', '墨西哥城国际', '墨西哥城1', '墨西哥城2'] * 2
            self.assertEqual([r[3] for r in list(workbook[app.SHEET_GO].values)[1:]], names)
            self.assertEqual([r[2] for r in list(workbook[app.SHEET_BACK].values)[1:]], names)
            self.assertEqual(workbook[app.SHEET_GO]['F2'].value, '=VLOOKUP(D2,机场列表!A:B,2,FALSE)')
            self.assertEqual(workbook[app.SHEET_BACK]['E2'].value, '=VLOOKUP(C2,机场列表!A:B,2,FALSE)')

    def test_cli_requires_confirmed_airline_before_reading_files(self):
        with redirect_stderr(StringIO()), self.assertRaises(SystemExit) as result:
            app.main(['merge', '--open', 'not-read.xlsx', '--temp', 'not-read.xlsx', '--effective', '2026-10-25', '--expiry', '2027-03-27'])
        self.assertEqual(result.exception.code, 2)

    def test_airline_scope_defaults_and_explicit_selection(self):
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            paths = [directory / 'open.xlsx', directory / 'temp.xlsx']
            for path in paths:
                workbook = openpyxl.Workbook()
                sheet = workbook.active
                sheet.append(['序号', '公司', '航站', '航站', '去程', '回程', '是否有台湾地区飞行员'])
                for i, code in enumerate(['CSN', 'CSG', ' csg ', 'CSGX'], 1):
                    sheet.append([i, code, '广州', '伦敦', code, '', '否'])
                workbook.save(path)
            output = directory / 'output.xlsx'
            report = app.run_merge(*paths, app.DEFAULT_TEMPLATE, output, date(2026, 10, 25), date(2027, 3, 27))
            self.assertEqual(report['airlines'], ['CSG'])
            self.assertEqual((report['raw_open_rows'], report['raw_temp_rows'], report['excluded_rows'], report['merged_rows']), (4, 4, 4, 4))
            workbook = openpyxl.load_workbook(output)
            self.assertEqual([r[1] for r in list(workbook[app.SHEET_MERGE].values)[1:]], ['CSG', 'csg', 'CSG', 'csg'])
            report = app.run_merge(*paths, app.DEFAULT_TEMPLATE, output, date(2026, 10, 25), date(2027, 3, 27), airlines=['CSN', 'CSGX'])
            self.assertEqual(report['merged_rows'], 4)
            for selection in [[], ['UNKNOWN']]:
                with self.assertRaises(ValueError):
                    app.run_merge(*paths, app.DEFAULT_TEMPLATE, output, date(2026, 10, 25), date(2027, 3, 27), airlines=selection)

    def test_unselected_bad_rows_do_not_block_selected_scope(self):
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            paths = [directory / 'open.xlsx', directory / 'temp.xlsx']
            for path in paths:
                workbook = openpyxl.Workbook()
                sheet = workbook.active
                sheet.append([1, 'CSG', '广州', '伦敦', 'G', 'B'])
                sheet.append([2, 'CSN', '广州', '', 'G', 'B'])
                sheet.append([3, 'CSN', '广州', '伦敦', datetime(2026, 10, 25), 'B'])
                workbook.save(path)
            output = directory / 'output.xlsx'
            report = app.run_merge(*paths, app.DEFAULT_TEMPLATE, output, date(2026, 10, 25), date(2027, 3, 27))
            self.assertEqual(report['merged_rows'], 2)
            with self.assertRaises(ValueError):
                app.run_merge(*paths, app.DEFAULT_TEMPLATE, output, date(2026, 10, 25), date(2027, 3, 27), airlines=['CSN'])
            workbook = openpyxl.load_workbook(paths[0])
            workbook.active.append([4, '', '广州', '伦敦', 'G', 'B'])
            workbook.save(paths[0])
            with self.assertRaises(ValueError):
                app.run_merge(*paths, app.DEFAULT_TEMPLATE, output, date(2026, 10, 25), date(2027, 3, 27))

    def test_header_alias_and_first_data_row(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["冬春批复"])
            sheet.append(["序号", "航空\n公司", "航站", "航站", "去程", "回程", "是否有台湾\n地区飞行员"])
            sheet.append([1, "CSG", "广州", "伦敦", "", "任意非空文本", "否"])
            sheet.append(["执行日期：2026年10月25日-2027年3月27日。"])
            workbook.save(path)
            rows = app.read_source_rows(path)
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]["source_row"], 3)
            self.assertEqual(rows[0]["company"], "CSG")
            workbook = openpyxl.Workbook()
            workbook.active.append([1, "CSG", "广州", "伦敦", "G", "B"])
            workbook.save(path)
            self.assertEqual(len(app.read_source_rows(path)), 1)

    def test_independent_directions_and_opaque_codes(self):
        route = "NIXAL-ZGGG-01R,SARIN-ZGGG-01,MORIT-ZGGG-01"
        base = {"company": "CSG", "station_a": "广州", "station_b": "伦敦", "tw": "否"}
        rows = [dict(base, go=g, back=b) for g, b in [("任意非空\n文本", route), (" \t", route), (route, "\n"), ("", "")]]
        go, back = app.build_legs(rows, {"伦敦": ["伦敦1", "伦敦2", "伦敦3"]})
        self.assertEqual([r["go"] for r in go], ["任意非空\n文本"] * 3 + [route] * 3)
        self.assertEqual([r["back"] for r in back], [route] * 6)
        self.assertEqual(back[0]["left"], "伦敦1")
        self.assertEqual(back[0]["right"], "广州")

    def test_export_preserves_od_and_has_no_empty_direction_rows(self):
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            paths = [directory / "open.xlsx", directory / "temporary.xlsx"]
            for i, path in enumerate(paths):
                workbook = openpyxl.Workbook()
                sheet = workbook.active
                sheet.append(["序号", "公司", "航站", "航站", "去程", "回程", "是否有台湾地区飞行员"])
                sheet.append([1, "CSG", "广州", "伦敦", "" if i == 0 else " ", "B" if i == 0 else "", "否"])
                workbook.save(path)
            output = directory / "result.xlsx"
            report = app.run_merge(*paths, app.DEFAULT_TEMPLATE, output, date(2026, 10, 25), date(2027, 3, 27))
            workbook = openpyxl.load_workbook(output)
            self.assertEqual(workbook[app.SHEET_MERGE].max_row, 3)
            self.assertEqual(workbook[app.SHEET_GO].max_row, 1)
            self.assertEqual(workbook[app.SHEET_BACK].max_row, 4)
            self.assertEqual(report["go_skipped_blank"], 2)
            self.assertEqual(report["both_blank"], 1)
            self.assertEqual(report["back_skipped_blank"], 1)


if __name__ == "__main__":
    unittest.main()
