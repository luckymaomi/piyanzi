# -*- coding: utf-8 -*-
"""按潜规则清洗后复算报务周期等，对照历史报告样例数字。"""
from __future__ import annotations

import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from statistics import mean, median

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(r"C:\Users\Administrator\Desktop\技术等级运行资格")
SP = BASE / "南货航飞行实力表20261023R2.xlsx"
TG = BASE / "技术等级导出.xlsx"
OQ = BASE / "运行资格导出.xlsx"
OUT = Path(__file__).resolve().parents[1] / "核对_报务周期与报告样例.txt"
CUTOFF = date(2026, 10, 23)  # 实力表版本日
REPORT = {
    "第一条": (538.56, 488.0, 99),
    "北美": (681.70, 681.5, 64),
    "欧洲": (582.35, 509.0, 85),
}
BAOWU_CODES = {"EAMA", "EEUO", "EWAS", "ESEA"}
TONGJI_YUAN = ("总队777", "上海", "新疆")


def to_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()[:10]
    for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def load_sheet(path, sheet=None):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    idx = {h: i for i, h in enumerate(headers) if h}
    return headers, idx, rows[1:]


def yuan_bucket(yuan: str) -> str:
    y = yuan or ""
    if "返聘" in y:
        return "返聘"
    if any(k in y for k in ("总队777", "上海", "新疆")):
        return "同机型划转"
    return "晋升降本"


def main():
    _, pi, prow = load_sheet(SP, "人员信息")
    _, ti, trow = load_sheet(TG)
    _, oi, orow = load_sheet(OQ)

    # 人员：副驾驶
    people = {}
    for r in prow:
        emp = str(r[pi["员工号"]] or "").strip()
        if not emp:
            continue
        tech = str(r[pi["技术信息"]] or "")
        if "副驾驶" not in tech:
            continue
        yuan = str(r[pi["原单位"]] or "")
        people[emp] = {
            "姓名": r[pi["姓名"]],
            "技术信息": tech,
            "原单位": yuan,
            "册": yuan_bucket(yuan),
        }

    fo_codes = {"FOA1", "FOA2", "FOB", "FOC", "FOD"}
    first = defaultdict(dict)  # emp -> code -> date
    for r in trow:
        emp = str(r[ti["员工号"]] or "").strip()
        if emp not in people:
            continue
        code = str(r[ti["技术等级代码"]] or "").strip().upper()
        if code not in fo_codes:
            continue
        ac = str(r[ti["机型"]] or "").strip()
        if ac != "777":
            continue
        d = to_date(r[ti["生效日期"]])
        if not d:
            continue
        prev = first[emp].get(code)
        if prev is None or d < prev:
            first[emp][code] = d

    def run_start(emp):
        return first[emp].get("FOA1")

    baowu_all = defaultdict(lambda: defaultdict(list))
    for r in orow:
        emp = str(r[oi["员工号"]] or "").strip()
        if emp not in people:
            continue
        code = str(r[oi["运行资格代码"]] or "").strip().upper()
        if code not in BAOWU_CODES:
            continue
        ac = str(r[oi["机型"]] or "").strip()
        if ac and ac != "777":
            continue
        d = to_date(r[oi["生效日期"]])
        if d:
            baowu_all[emp][code].append(d)

    def cycles_for_v2(emps, label):
        lines = [f"【{label}】人数={len(emps)}"]
        specs = [("第一条", None), ("北美", "EAMA"), ("欧洲", "EEUO")]
        for name, code in specs:
            days_list = []
            marked_early = 0
            for emp in emps:
                s = run_start(emp)
                if not s:
                    continue
                if code is None:
                    pool = []
                    for c in BAOWU_CODES:
                        pool.extend(baowu_all[emp].get(c, []))
                else:
                    pool = list(baowu_all[emp].get(code, []))
                if not pool:
                    continue
                if any(d < s for d in pool):
                    marked_early += 1
                after = [d for d in pool if d >= s]
                if not after:
                    continue
                e = min(after)
                days_list.append((e - s).days)
            if days_list:
                avg = mean(days_list)
                med = float(median(days_list))
                lines.append(f"  {name}: 平均={avg:.2f} 中位={med} n={len(days_list)} (早于运行已标记人数触碰={marked_early})")
                if name in REPORT:
                    ra, rm, rn = REPORT[name]
                    lines.append(
                        f"    对照报告: {ra} / {rm} / n={rn} → Δ平均={avg - ra:+.2f} Δ中位={med - rm:+.1f} Δn={len(days_list) - rn:+d}"
                    )
            else:
                lines.append(f"  {name}: 无有效样本")
        no_any = []
        for emp in emps:
            s = run_start(emp)
            if not s:
                continue
            pool = []
            for c in BAOWU_CODES:
                pool.extend(baowu_all[emp].get(c, []))
            after = [d for d in pool if d >= s]
            if after:
                continue
            no_any.append((CUTOFF - s).days)
        if no_any:
            lines.append(
                f"  无报务盘: n={len(no_any)} 已运行平均={mean(no_any):.1f} 中位={median(no_any)} 最长={max(no_any)} 最短={min(no_any)}"
            )
        foc = [(first[emp]["FOC"] - run_start(emp)).days for emp in emps if run_start(emp) and first[emp].get("FOC") and first[emp]["FOC"] >= run_start(emp)]
        if foc:
            lines.append(f"  运行→FOC生效: 平均={mean(foc):.2f} 中位={median(foc)} n={len(foc)}")
        return "\n".join(lines)

    all_fo = list(people)
    jinsheng = [e for e, p in people.items() if p["册"] == "晋升降本"]
    huazhuan = [e for e, p in people.items() if p["册"] == "同机型划转"]
    fanpin = [e for e, p in people.items() if p["册"] == "返聘"]
    # 有 FOA1 的子集
    with_a1 = [e for e in all_fo if run_start(e)]
    jinsheng_a1 = [e for e in jinsheng if run_start(e)]

    # 结构快照
    from collections import Counter

    tech_c = Counter()
    for p in people.values():
        t = p["技术信息"]
        for tag in ("划转副驾驶", "A1", "A2", "B类", "C类", "D类", "E类"):
            if tag in t.replace(" ", ""):
                pass
        # rough
        if "划转副驾驶" in t:
            tech_c["划转副驾驶→A1口径"] += 1
        elif "D类" in t or ":D" in t or "D类副驾驶" in t:
            tech_c["D"] += 1
        elif "C类" in t:
            tech_c["C"] += 1
        elif "B类" in t:
            tech_c["B"] += 1
        elif "A2" in t:
            tech_c["A2"] += 1
        elif "A1" in t:
            tech_c["A1"] += 1
        elif "E类" in t:
            tech_c["E"] += 1
        else:
            tech_c[t] += 1

    parts = []
    parts.append(f"截止日={CUTOFF} 实力表副驾驶={len(people)} 有FOA1={len(with_a1)}")
    parts.append(f"分册: 晋升降本={len(jinsheng)} 同机型划转={len(huazhuan)} 返聘={len(fanpin)}")
    parts.append(f"技术信息粗分: {dict(tech_c)}")
    parts.append("")
    parts.append("===== 报务周期复算（对照报告样例 538.56/488/99；681.70/681.5/64；582.35/509/85）=====")
    parts.append(cycles_for_v2(with_a1, "全副驾驶·有FOA1（接近旧报告全量口径）"))
    parts.append("")
    parts.append(cycles_for_v2(jinsheng_a1, "仅晋升降本·有FOA1（潜规则成长主盘）"))
    parts.append("")
    parts.append(cycles_for_v2(huazhuan, "仅同机型划转（有无A1都进筛选，内部再看起点）"))

    text = "\n".join(parts)
    OUT.write_text(text, encoding="utf-8")
    print(text)
    print("\n写入", OUT)


if __name__ == "__main__":
    main()
