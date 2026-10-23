# -*- coding: utf-8 -*-
"""One-shot: infer validity length and scheduling lead from 南货航培训总表."""
from __future__ import annotations

from collections import Counter
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl.worksheet.datavalidation import DataValidation

_orig = DataValidation.__init__


def _init(self, *a, **k):
    k.pop("id", None)
    return _orig(self, *a, **k)


DataValidation.__init__ = _init

import openpyxl  # noqa: E402

PATH = Path(r"C:\Users\Administrator\Desktop\南货航2026年培训表 (1).xlsx")


def parse_d(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def add_months(d: date, m: int) -> date:
    y = d.year + (d.month - 1 + m) // 12
    mo = (d.month - 1 + m) % 12 + 1
    dim = [31, 29 if y % 4 == 0 and (y % 100 != 0 or y % 400 == 0) else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31][mo - 1]
    return date(y, mo, min(d.day, dim))


def month_end(d: date) -> date:
    return add_months(date(d.year, d.month, 1), 1) - timedelta(days=1)


def summ(arr, label):
    if not arr:
        print(f"  {label}: n=0")
        return
    s = sorted(arr)

    def p(q):
        return s[min(len(s) - 1, int(len(s) * q))]

    print(
        f"  {label}: n={len(s)} min={s[0]} p10={p(0.1)} p25={p(0.25)} "
        f"med={p(0.5)} p75={p(0.75)} p90={p(0.9)} max={s[-1]}"
    )


RULES = {
    "应急训练": dict(months=24, round="month_end", windowed=True),
    "危险品": dict(months=24, round="day", windowed=True),  # 2年
    "航空安保": dict(months=24, round="month_end", windowed=False),
    "TSA": dict(months=12, round="month_end", windowed=False),
    "疲劳管理": dict(months=24, round="month_end", windowed=False),
    "飞行作风": dict(months=24, round="month_end", windowed=False),
    "英语能力": dict(months=36, round="day", windowed=True),
    "汉语能力": dict(months=72, round="day", windowed=True),
}


def main():
    wb = openpyxl.load_workbook(PATH, data_only=True)
    people = next(wb[n] for n in wb.sheetnames if "人员信息" in n)
    ph = [c.value for c in people[1]]
    proj_cols = {
        str(h).strip(): i
        for i, h in enumerate(ph)
        if h and str(h).strip() in RULES
    }
    id_i = next(i for i, h in enumerate(ph) if h and "员工号" in str(h))
    people_by_id = {}
    for row in people.iter_rows(min_row=2, values_only=True):
        eid = row[id_i]
        if eid is None:
            continue
        key = str(int(eid)) if isinstance(eid, (int, float)) else str(eid).strip()
        rec = {}
        for p, ci in proj_cols.items():
            v = row[ci]
            rec[p] = v.date() if isinstance(v, datetime) else (v if isinstance(v, date) else None)
        people_by_id[key] = rec

    print("people", len(people_by_id))

    for proj, rule in RULES.items():
        if proj not in wb.sheetnames:
            print(f"==== {proj} NO SHEET ====")
            continue
        ws = wb[proj]
        hdr = [c.value for c in ws[1]]
        if "培训开始日期" not in hdr:
            print(f"==== {proj} no start col ====")
            continue
        si, ei = hdr.index("培训开始日期"), hdr.index("员工号")
        note_i = hdr.index("备注") if "备注" in hdr else None
        sei = hdr.index("有效期") if "有效期" in hdr else None

        validity_days = []
        month_diffs = []
        lead_est = []
        train_months = Counter()
        nominal_match = 0
        rows_n = 0

        for row in ws.iter_rows(min_row=2, values_only=True):
            if note_i is not None and row[note_i] and "取消" in str(row[note_i]):
                continue
            start = parse_d(row[si])
            if not start:
                continue
            rows_n += 1
            train_months[start.strftime("%Y-%m")] += 1
            exp = parse_d(row[sei]) if sei is not None else None
            eid = row[ei]
            key = str(int(eid)) if isinstance(eid, (int, float)) else str(eid or "").strip()
            if not exp and key in people_by_id:
                exp = people_by_id[key].get(proj)
            if not exp:
                continue

            vd = (exp - start).days
            validity_days.append(vd)
            month_diffs.append((exp.year - start.year) * 12 + (exp.month - start.month))

            if rule["round"] == "month_end":
                expected = month_end(add_months(start, rule["months"]))
            else:
                expected = add_months(start, rule["months"])
            if abs((exp - expected).days) <= 3:
                nominal_match += 1

            # 推定旧到期：命中窗口时 new≈old+validity → old=new-validity
            if rule["round"] == "day":
                old = add_months(exp, -rule["months"])
            else:
                old = month_end(add_months(date(exp.year, exp.month, 1), -rule["months"]))
            lead_est.append((old - start).days)

        print(f"==== {proj} ====")
        print(f"  rule months={rule['months']} round={rule['round']} windowed={rule['windowed']}")
        print(f"  rows={rows_n} with_exp={len(validity_days)} exact_recalc_from_start={nominal_match}")
        print("  train_months", dict(sorted(train_months.items())))
        summ(validity_days, "newExpiry-start days")
        summ(month_diffs, "newExpiry-start monthDiff")
        summ(lead_est, "est lead days (old=new-validity) - start")
        sane = [x for x in lead_est if -15 <= x <= 100]
        summ(sane, "est lead sane[-15,100]")
        print(
            f"  lead buckets: "
            f"<0={sum(1 for x in lead_est if x < 0)} "
            f"0-30={sum(1 for x in lead_est if 0 <= x <= 30)} "
            f"31-60={sum(1 for x in lead_est if 31 <= x <= 60)} "
            f"61-90={sum(1 for x in lead_est if 61 <= x <= 90)} "
            f"91-120={sum(1 for x in lead_est if 91 <= x <= 120)} "
            f">120={sum(1 for x in lead_est if x > 120)}"
        )

    print("\n==== people expiry pressure (2026-2028) ====")
    for proj in ["TSA", "航空安保", "应急训练", "危险品", "飞行作风", "英语能力", "汉语能力", "疲劳管理"]:
        c = Counter()
        for rec in people_by_id.values():
            e = rec.get(proj)
            if e and 2026 <= e.year <= 2028:
                c[e.strftime("%Y-%m")] += 1
        print(proj, "top", c.most_common(6), "total", sum(c.values()))

    wb.close()


if __name__ == "__main__":
    main()
