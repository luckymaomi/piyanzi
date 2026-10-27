# -*- coding: utf-8 -*-
"""副驾驶 A2→B 超过 1 年 9 个月名单。Excel + Markdown。"""
from __future__ import annotations

import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
import openpyxl.utils

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(r"C:\Users\Administrator\Desktop\技术等级运行资格")
SP = BASE / "南货航飞行实力表20261023R2.xlsx"
TG = BASE / "技术等级导出.xlsx"
ROOT = Path(__file__).resolve().parents[1]
CUTOFF = date(2026, 10, 23)
# 1年9个月：365 + 9×30 = 635 天（口径写死，避免歧义）
LIMIT = 365 + 9 * 30
OUT_X = ROOT / "副驾驶_A2转B超过1年9个月名单_20261023.xlsx"
OUT_MD = ROOT / "副驾驶_A2转B超过1年9个月_说明_20261023.md"


def to_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()[:10]
    for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def load(path, sheet=None):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    h = [str(x).strip() if x is not None else "" for x in rows[0]]
    return {x: i for i, x in enumerate(h) if x}, rows[1:]


def yuan_bucket(yuan: str) -> str:
    y = yuan or ""
    if "返聘" in y:
        return "返聘"
    if any(k in y for k in ("总队777", "上海", "新疆")):
        return "同机型划转"
    return "晋升降本"


def style_header(ws, n):
    fill = PatternFill("solid", fgColor="D9E2F3")
    for c in range(1, n + 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    if ws.max_row >= 1:
        ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(n)}{max(ws.max_row, 1)}"


def main():
    pi, prow = load(SP, "人员信息")
    ti, trow = load(TG)

    people = {}
    for r in prow:
        emp = str(r[pi["员工号"]] or "").strip()
        if not emp:
            continue
        tech = str(r[pi["技术信息"]] or "")
        if "副驾驶" not in tech:
            continue
        people[emp] = dict(
            员工号=emp,
            姓名=r[pi["姓名"]],
            分部=r[pi["分部"]],
            技术信息=tech,
            原单位=r[pi["原单位"]],
            册=yuan_bucket(str(r[pi["原单位"]] or "")),
        )

    first = defaultdict(dict)
    for r in trow:
        emp = str(r[ti["员工号"]] or "").strip()
        if emp not in people:
            continue
        code = str(r[ti["技术等级代码"]] or "").strip().upper()
        if code not in ("FOA2", "FOB"):
            continue
        if str(r[ti["机型"]] or "").strip() != "777":
            continue
        d = to_date(r[ti["生效日期"]])
        if not d:
            continue
        if code not in first[emp] or d < first[emp][code]:
            first[emp][code] = d

    done, stuck = [], []
    for emp, p in people.items():
        a2 = first[emp].get("FOA2")
        b = first[emp].get("FOB")
        if a2 and b and b >= a2:
            days = (b - a2).days
            if days > LIMIT:
                done.append({**p, "FOA2首次生效": a2, "FOB首次生效": b, "间隔天": days, "类型": "已完成且间隔>1年9个月"})
        elif a2 and not b:
            days = (CUTOFF - a2).days
            if days > LIMIT:
                stuck.append({**p, "FOA2首次生效": a2, "已耗天": days, "类型": "仍无B且自A2已超1年9个月"})

    wb = Workbook()
    ws = wb.active
    ws.title = "00_口径"
    for a, b in [
        ("任务", "副驾驶 A2转B 超过 1年9个月"),
        ("实力表", str(SP)),
        ("技术等级", str(TG)),
        ("截止日", str(CUTOFF)),
        ("机型", "777"),
        ("节点", "FOA2 / FOB 首次生效"),
        ("阈值天数", f"{LIMIT}（365+9×30，即1年9个月）"),
        ("名单范围", "人员信息·副驾驶"),
        ("已完成超期", len(done)),
        ("在途超期", len(stuck)),
        ("副驾驶池", len(people)),
    ]:
        ws.append([a, b])
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 90

    def sheet(title, rows, headers):
        w = wb.create_sheet(title)
        w.append(headers)
        key = "间隔天" if "间隔天" in headers else "已耗天"
        for r in sorted(rows, key=lambda x: (-(x.get(key) or 0), x["员工号"])):
            w.append([r.get(h) for h in headers])
        if rows:
            style_header(w, len(headers))

    sheet(
        "01_已完成间隔超1年9个月",
        done,
        ["员工号", "姓名", "分部", "册", "技术信息", "原单位", "FOA2首次生效", "FOB首次生效", "间隔天", "类型"],
    )
    sheet(
        "02_在途超1年9个月未到B",
        stuck,
        ["员工号", "姓名", "分部", "册", "技术信息", "原单位", "FOA2首次生效", "已耗天", "类型"],
    )
    ws = wb.create_sheet("03_汇总")
    ws.append(["类别", "人数"])
    ws.append(["已完成且间隔>1年9个月", len(done)])
    ws.append(["在途超1年9个月未到B", len(stuck)])
    ws.append(["副驾驶池", len(people)])
    wb.save(OUT_X)

    md = f"""# 副驾驶 A2转B 超过 1年9个月

## 0. 本轮做了什么

| 项 | 值 |
| --- | --- |
| 截止日 | {CUTOFF} |
| 实力表 | `{SP}` |
| 技术等级 | `{TG}` |
| 阈值 | 1年9个月 = **{LIMIT}** 天（365+9×30） |

---

## 1. 清洗后是什么样

副驾驶池 **{len(people)}** 人；777 上 FOA2/FOB 取首次生效。

---

## 2. 按什么口径统计了什么

1. 起点 FOA2 首次生效，终点 FOB 首次生效。  
2. **已完成超期**：间隔 **> {LIMIT}** 天 → **{len(done)}** 人。  
3. **在途超期**：有 FOA2 无 FOB，截止日已耗 **> {LIMIT}** 天 → **{len(stuck)}** 人。

---

## 3. 结论

1. 已转 B 且 A2→B 超过 1年9个月：**{len(done)}** 人（Excel 01）。  
2. 仍未到 B 且自 A2 已超过 1年9个月：**{len(stuck)}** 人（Excel 02）。

---

## 4. 附件

1. `{OUT_X}`  
2. 脚本：`scripts/list_a2b_over_1y9m.py`  
"""
    OUT_MD.write_text(md, encoding="utf-8")
    print(md)


if __name__ == "__main__":
    main()
